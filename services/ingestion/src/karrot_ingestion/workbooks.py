from __future__ import annotations

import hashlib
import json
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook

from .normalization import abn_text, json_value, postcode_text, row_hash


MANIFEST_PATH = Path("data/manifests/source-manifest.json")

PROVIDER_DATA_SHEETS = (
    "Provider Details",
    "Residential Care Home Details",
    "Category & Service Type Details",
    "Category & LGA Details",
    "Regulatory Notices",
)
STAR_DATA_SHEETS = ("Star Ratings", "Detailed data")
CARE_DATA_SHEETS = ("Jul-Sep 2025", "Oct-Dec 2025", "Jan-Mar 2026")

CARE_PERIODS = {
    "Jul-Sep 2025": (date(2025, 7, 1), date(2025, 9, 30)),
    "Oct-Dec 2025": (date(2025, 10, 1), date(2025, 12, 31)),
    "Jan-Mar 2026": (date(2026, 1, 1), date(2026, 3, 31)),
}


@dataclass(frozen=True)
class SourceRow:
    sheet_name: str
    row_number: int
    values: dict[str, Any]

    @property
    def hash(self) -> str:
        return row_hash(self.values)

    @property
    def json(self) -> dict[str, Any]:
        return {key: json_value(value) for key, value in self.values.items()}


def file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def load_manifest(data_dir: Path) -> dict[str, dict[str, Any]]:
    path = data_dir / MANIFEST_PATH
    payload = json.loads(path.read_text(encoding="utf-8"))
    return {item["dataset_code"]: item for item in payload["datasets"]}


def validate_manifest_files(data_dir: Path, manifest: dict[str, dict[str, Any]]) -> None:
    errors = []
    for item in manifest.values():
        path = data_dir / item["relative_path"]
        if not path.exists():
            errors.append(f"Missing source file: {path}")
            continue
        actual = file_sha256(path)
        if actual != item["sha256"]:
            errors.append(
                f"Checksum mismatch for {path.name}: expected {item['sha256']}, got {actual}"
            )
    if errors:
        raise ValueError("\n".join(errors))


def _sheet_rows(path: Path, sheet_name: str, header_row: int) -> list[SourceRow]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook[sheet_name]
    headers = [cell.value for cell in sheet[header_row]]
    last_header = max(
        index for index, value in enumerate(headers, start=1) if value not in (None, "")
    )
    clean_headers = [str(value).strip() for value in headers[:last_header]]
    result = []
    for row_number, cells in enumerate(
        sheet.iter_rows(min_row=header_row + 1, max_col=last_header),
        start=header_row + 1,
    ):
        values = [cell.value for cell in cells]
        if not any(value not in (None, "") for value in values):
            continue
        payload = dict(zip(clean_headers, values))
        if "ABN" in payload:
            payload["ABN"] = abn_text(payload["ABN"])
        for key in ("Postcode", "Residential Care Home Postcode"):
            if key in payload:
                payload[key] = postcode_text(payload[key])
        result.append(SourceRow(sheet_name, row_number, payload))
    workbook.close()
    return result


def read_provider_register(data_dir: Path) -> dict[str, list[SourceRow]]:
    path = data_dir / load_manifest(data_dir)["acqsc_provider_register"]["relative_path"]
    return {name: _sheet_rows(path, name, 1) for name in PROVIDER_DATA_SHEETS}


def read_star_ratings(data_dir: Path) -> dict[str, list[SourceRow]]:
    path = data_dir / load_manifest(data_dir)["star_ratings"]["relative_path"]
    return {name: _sheet_rows(path, name, 1) for name in STAR_DATA_SHEETS}


def read_care_minutes(data_dir: Path) -> dict[str, list[SourceRow]]:
    path = data_dir / load_manifest(data_dir)["care_minutes"]["relative_path"]
    return {name: _sheet_rows(path, name, 2) for name in CARE_DATA_SHEETS}


def all_rows(sheets: dict[str, list[SourceRow]]) -> Iterable[SourceRow]:
    for rows in sheets.values():
        yield from rows


def parse_date(value: Any) -> date | None:
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return date.fromisoformat(str(value)[:10])


def column_starting_with(payload: dict[str, Any], prefix: str) -> Any:
    for key, value in payload.items():
        if key.startswith(prefix):
            return value
    raise KeyError(f"No column starting with {prefix!r}")
